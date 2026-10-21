import json, sys
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn

S, OUT = sys.argv[1], sys.argv[2]
dims = json.load(open(f'{S}/lh-dims.json'))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(18)
sec.top_margin = Mm(dims['headMm'] + 4)
sec.bottom_margin = Mm(dims['footMm'] + 6)
sec.header_distance = sec.footer_distance = Mm(0)

def full_bleed(par, img, height_mm):
    pf = par.paragraph_format
    pf.left_indent = pf.right_indent = Mm(-18)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    par.add_run().add_picture(img, width=Mm(210))

full_bleed(sec.header.paragraphs[0], f'{S}/lh-header.png', dims['headMm'])
full_bleed(sec.footer.paragraphs[0], f'{S}/lh-footer.png', dims['footMm'])

# Body font: Latin + Bangla (complex script) faces that ship with Windows and Office
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
rpr = style.element.get_or_add_rPr()
fonts = rpr.find(qn('w:rFonts'))
if fonts is None:
    fonts = rpr.makeelement(qn('w:rFonts'), {}); rpr.append(fonts)
fonts.set(qn('w:cs'), 'Nirmala UI'); fonts.set(qn('w:eastAsia'), 'Nirmala UI')
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.25

grey = RGBColor(0x4f, 0x61, 0x58)
p = doc.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Mm(80), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
p.paragraph_format.tab_stops.add_tab_stop(Mm(118), WD_TAB_ALIGNMENT.LEFT)
p.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
for text in ['স্মারক নং / Ref: ', '\t', '\tতারিখ / Date: ', '\t']:
    r = p.add_run(text); r.font.color.rgb = grey; r.font.size = Pt(10)
doc.add_paragraph()
for line in ['প্রাপক / To:', '', 'বিষয় / Subject:', '', 'জনাব / Dear Sir or Madam,', '', '', '', 'বিনীত / Sincerely,', '', '', '', 'বারিজাহ সলিউশনস / Bareezah Solutions']:
    doc.add_paragraph(line)
doc.core_properties.title = 'Bareezah Solutions letterhead'
doc.core_properties.author = 'Bareezah Solutions'
doc.save(OUT)
print('saved', OUT)
